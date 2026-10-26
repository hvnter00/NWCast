from __future__ import annotations
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd

from prediction_io import (
    PRICE_COLUMNS, _ALIASES, _COLUMN_NAMES, _blank, _date_value,
    _heading, _number_value, make_template, read_upload,
)


VARIABLES = tuple(PRICE_COLUMNS)
_FILES = {
    "coal": ("Coal.csv", "NEWCASTLE.csv"),
    "wti": ("WTI OIL.csv", "WTI.csv"),
    "gas": ("HenryHub.csv", "Gas.csv"),
    "dxy": ("DXY.csv",),
}
_ORIGINAL_FORMAT = {
    "coal": ("Tanggal", "Terakhir", "%d %m %Y", ","),
    "wti": ("observation_date", "DCOILWTICO", "%d %m %Y", "."),
    "gas": ("Date", "Henry Hub", "%b %d, %Y", ","),
    "dxy": ("Tanggal", "Terakhir", "%d %m %Y", ","),
}
_KNOWN_DXY_SATURDAYS = {pd.Timestamp("2025-05-17"), pd.Timestamp("2025-12-06")}


def _validate_raw(raw: pd.DataFrame) -> pd.DataFrame:
    if not isinstance(raw, pd.DataFrame) or not isinstance(raw.index, pd.DatetimeIndex):
        raise ValueError("Indeks tanggal data tidak valid.")
    if raw.empty or raw.index.hasnans or raw.index.has_duplicates:
        raise ValueError("Data kosong atau tanggal duplikat.")
    if raw.index.tz is not None or (raw.index != raw.index.normalize()).any():
        raise ValueError("Gunakan tanggal tanpa jam/zona waktu.")
    if (raw.index.dayofweek >= 5).any():
        raise ValueError("Tanggal harus Senin–Jumat.")
    if raw.columns.duplicated().any() or not set(VARIABLES).issubset(raw.columns):
        raise ValueError("Kolom data belum lengkap.")
    try:
        result = raw.loc[:, list(VARIABLES)].astype(float).sort_index()
    except (ValueError, TypeError) as error:
        raise ValueError("Harga harus berupa angka.") from error
    if np.isinf(result.to_numpy()).any():
        raise ValueError("Harga tidak boleh infinity.")
    result.index.name = "Date"
    return result


def _read_original(path: Path, variable: str) -> pd.Series:
    payload = path.read_bytes()
    if not payload.strip():
        raise ValueError(f"CSV kosong: {path.name}.")
    try:
        header = payload.decode("utf-8-sig").splitlines()[0]
        raw = pd.read_csv(BytesIO(payload), sep=";" if ";" in header else ",",
                          dtype=str, keep_default_na=False, encoding="utf-8-sig")
        raw.columns = raw.columns.str.strip()
        spec = ("Date", "Price", "%Y-%m-%d", ".") if set(raw.columns) == {"Date", "Price"} else _ORIGINAL_FORMAT[variable]
        date_column, price_column, date_format, decimal = spec
        date = pd.to_datetime(raw[date_column].str.strip(), format=date_format, errors="coerce")
        text = raw[price_column].str.strip()
        if decimal == ",":
            text = text.str.replace(",", ".", regex=False)
        prices = pd.to_numeric(text.mask(text.eq("")), errors="coerce")
        if date.isna().any() or (prices.isna() & text.ne("")).any() or np.isinf(prices).any():
            raise ValueError("Tanggal/harga tidak valid.")
        frame = pd.DataFrame({"Date": date, "Price": prices})
    except (ValueError, KeyError, UnicodeError, pd.errors.ParserError) as error:
        raise ValueError(f"Periksa CSV {path.name}.") from error

    if frame.groupby("Date")["Price"].nunique(dropna=True).gt(1).any():
        raise ValueError(f"Harga duplikat berbeda: {path.name}.")
    frame = frame.sort_values("Date").groupby("Date", as_index=False)["Price"].first()
    weekend = frame.Date.dt.dayofweek.ge(5)
    for row in frame.loc[weekend].itertuples(index=False):
        previous = frame.loc[frame.Date.eq(row.Date - pd.offsets.BDay(1)), "Price"]
        known_duplicate = (
            variable == "dxy" and row.Date in _KNOWN_DXY_SATURDAYS
            and len(previous) == 1 and pd.notna(row.Price) and row.Price == previous.iloc[0]
        )
        if not known_duplicate:
            raise ValueError(f"Observasi akhir pekan: {path.name}, {row.Date:%Y-%m-%d}.")
    return frame.loc[~weekend].set_index("Date")["Price"].rename(variable)


def load_historical(base_dir: str | Path) -> pd.DataFrame:
    """Load published raw data, or bootstrap the notebook's four original CSVs."""
    from model_registry import active_dataset_path

    base = Path(base_dir)
    published = active_dataset_path(base)
    if published is not None:
        try:
            frame = pd.read_csv(published, dtype={column: float for column in VARIABLES},
                                float_precision="round_trip")
            frame["Date"] = pd.to_datetime(frame["Date"], format="%Y-%m-%d", errors="raise")
            return _validate_raw(frame.set_index("Date"))
        except (OSError, ValueError, KeyError, pd.errors.ParserError) as error:
            raise ValueError("Dataset versi aktif tidak dapat dibaca.") from error

    series = []
    for variable, filenames in _FILES.items():
        found = [base / "Data" / "Historical" / name for name in filenames
                 if (base / "Data" / "Historical" / name).is_file()]
        if len(found) != 1:
            raise ValueError(f"CSV historis {variable.upper()} tidak ditemukan atau ganda.")
        series.append(_read_original(found[0], variable))
    return _validate_raw(pd.concat(series, axis=1).sort_index())


def parse_new_data(payload: bytes, filename: str) -> pd.DataFrame:
    """Parse combined observations; blank cells never erase existing prices."""
    frame = read_upload(payload, filename)
    if frame.empty:
        raise ValueError("File tidak berisi data.")
    selected = {}
    for column in frame.columns:
        canonical = _ALIASES.get(_heading(column))
        if canonical is None:
            if frame[column].map(_blank).all():
                continue
            raise ValueError(f"Kolom tidak dikenal: {column}.")
        if canonical in selected:
            raise ValueError("Ada nama kolom ganda.")
        selected[canonical] = column
    required = ["date", *VARIABLES]
    if any(key not in selected for key in required):
        raise ValueError("Gunakan kolom Date, Coal, Oil, Gas, DXY.")
    parsed = []
    for row_number, (_, source) in enumerate(frame.iterrows(), start=1):
        if all(_blank(source[selected[key]]) for key in required):
            continue
        try:
            timestamp = _date_value(source[selected["date"]])
        except (ValueError, TypeError, OverflowError) as error:
            raise ValueError(f"Baris {row_number}: tanggal tidak valid.") from error
        if timestamp.dayofweek >= 5:
            raise ValueError(f"Baris {row_number}: tanggal harus Senin–Jumat.")
        values = {}
        for column in VARIABLES:
            item = source[selected[column]]
            if _blank(item):
                values[column] = np.nan
                continue
            try:
                value = _number_value(item)
            except (ValueError, TypeError, OverflowError) as error:
                raise ValueError(f"Baris {row_number}: harga {_COLUMN_NAMES[column]} tidak valid.") from error
            if value <= 0 and column != "wti":
                raise ValueError(f"Baris {row_number}: harga {_COLUMN_NAMES[column]} harus positif.")
            values[column] = value
        if any(pd.notna(value) for value in values.values()):
            parsed.append({"Date": timestamp, **values})
    if not parsed:
        raise ValueError("Isi minimal satu harga.")
    result = pd.DataFrame(parsed).set_index("Date").sort_index()
    if result.index.has_duplicates:
        if result.groupby(level=0).nunique(dropna=True).gt(1).any().any():
            raise ValueError("Ada harga berbeda pada tanggal duplikat.")
        result = result.groupby(level=0).first()
    return _validate_raw(result)


def merge_data(historical: pd.DataFrame, new: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Apply provided values only; keep the published input untouched until success."""
    old = _validate_raw(historical)
    incoming = _validate_raw(new)
    full_index = old.index.union(incoming.index).sort_values()
    before = old.reindex(full_index)
    uploaded = incoming.reindex(full_index)
    provided = uploaded.notna()
    value_added = provided & before.isna()
    value_changed = provided & before.notna() & uploaded.ne(before)
    merged = before.where(~provided, uploaded)
    dates_added = incoming.index.difference(old.index)
    updated = (value_changed | value_added).any(axis=1) & full_index.isin(old.index)
    summary = {
        "uploaded_rows": len(incoming),
        "dates_added": len(dates_added),
        "dates_updated": int(updated.sum()),
        "values_added": int(value_added.to_numpy().sum()),
        "values_changed": int(value_changed.to_numpy().sum()),
        "date_start": merged.index.min().date().isoformat(),
        "date_end": merged.index.max().date().isoformat(),
        "rows_before": len(old),
        "rows_after": len(merged),
    }
    return _validate_raw(merged), summary


def prepare_training_prices(raw: pd.DataFrame) -> pd.DataFrame:
    """Reproduce common overlap and forward fill from the original notebook."""
    frame = _validate_raw(raw)
    first = [frame[column].first_valid_index() for column in VARIABLES]
    last = [frame[column].last_valid_index() for column in VARIABLES]
    if any(value is None for value in first + last):
        raise ValueError("Setiap variabel perlu minimal satu harga.")
    start, end = max(first), min(last)
    if start > end:
        raise ValueError("Rentang tanggal antarvariabel tidak beririsan.")
    full_grid = pd.bdate_range(min(first), end, name="Date")
    common_grid = pd.bdate_range(start, end, name="Date")
    prices = frame.reindex(full_grid).ffill().loc[common_grid, list(VARIABLES)].copy()
    if not np.isfinite(prices.to_numpy()).all():
        raise ValueError("Harga awal setiap variabel belum lengkap.")
    return prices


def make_admin_template(format: str = "xlsx") -> bytes:
    payload = make_template(format=format)
    if format.lower().lstrip(".") == "csv":
        return payload
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(payload))
    guide = workbook["Panduan"]
    guide.delete_rows(1, guide.max_row)
    lines = [
        "Panduan upload data Admin",
        "Isi sheet Input dengan Date, Coal, Oil, Gas, DXY. Boleh satu atau banyak tanggal.",
        "Tanggal harus Senin–Jumat; gunakan tanggal Excel, YYYY-MM-DD, atau DD/MM/YYYY.",
        "Harga kosong diperbolehkan. Nilai lama tidak dihapus oleh sel kosong.",
        "Harga terisi pada tanggal lama akan mengganti nilai lama saat retraining berhasil.",
        "Coal, Gas, DXY harus positif. Oil historis nonpositif dipertahankan sesuai sumber data.",
        "Harga pada 31 tanggal terakhir dataset gabungan harus positif untuk fitur log.",
        "Tanggal tanpa observasi akan memakai harga sebelumnya saat preprocessing.",
        "Baris kosong diabaikan; duplikat dengan harga berbeda ditolak.",
        "Data gabungan dan model versi baru disimpan setelah retraining berhasil.",
    ]
    for line in lines:
        guide.append([line])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
